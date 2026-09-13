"""Regenerate the binary golden scripts: ``veri_scenes.pdf`` and ``ardha_song.docx``.

Run ``uv run python tests/golden/build_goldens.py``. The PDF is written by hand (no PDF
writer is a project dependency) as a two-page Courier document pdfplumber can read; the
DOCX is a plain Normal-style Word document (python-docx) with a page break before scene 2.
Both fixtures are original test material.
"""
from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent

# --- veri_scenes.pdf --------------------------------------------------------------------

X = {"pageno": 540, "action": 108, "slug": 108, "trans": 400, "char": 266, "paren": 223, "dlg": 180, "furniture": 440}
LEADING = 14
TOP_Y = 720

VERI_PAGES: list[list[tuple[str, str]]] = [
    [
        ("pageno", "1."),
        ("blank", ""),
        ("action", "VERI - three scenes (golden test fixture)"),
        ("blank", ""),
        ("trans", "FADE IN:"),
        ("blank", ""),
        ("slug", "EXT. THRISSUR POORAM GROUND - DAY"),
        ("blank", ""),
        ("action", "Elephants. Umbrellas. A CROWD of thousands under a white sun."),
        ("action", "VELU (40, ex-mahout, a limp) pushes through with a"),
        ("action", "STEEL TIFFIN CARRIER held over his head."),
        ("blank", ""),
        ("char", "VELU"),
        ("dlg", "Enthoottu thirakka ithu. Vazhi tharoo,"),
        ("dlg", "vazhi tharoo!"),
        ("blank", ""),
        ("action", "A DRUMMER on the temple platform catches his eye. Nods once."),
        ("blank", ""),
        ("slug", "INT. TEMPLE OFFICE - CONTINUOUS"),
        ("blank", ""),
        ("action", "A ceiling fan and a locked cupboard. The TEMPLE SECRETARY (60)"),
        ("action", "counts donation slips."),
        ("blank", ""),
        ("char", "VELU"),
        ("dlg", "Secretary sir, aana onnu nokkanam. Pani undu."),
        ("blank", ""),
        ("char", "TEMPLE SECRETARY"),
        ("dlg", "Pooram kazhinjittu mathi. Ippo pattilla."),
        ("blank", ""),
        ("action", "Velu sets the tiffin carrier on the desk. The secretary stops"),
        ("action", "counting."),
        ("blank", ""),
        ("furniture", "(CONTINUED)"),
    ],
    [
        ("pageno", "2."),
        ("blank", ""),
        ("furniture", "CONTINUED:"),
        ("blank", ""),
        ("slug", "EXT. ELEPHANT SHED - NIGHT"),
        ("blank", ""),
        ("action", "Chains. Torchlight. The elephant RAMANKUTTY shifts his weight,"),
        ("action", "one ear flapping. Velu opens the tiffin carrier: jaggery, and a"),
        ("action", "folded court summons."),
        ("blank", ""),
        ("char", "VELU"),
        ("paren", "(to the elephant)"),
        ("dlg", "Nee mathram ariyum. Naale ellaarum ariyum."),
        ("blank", ""),
        ("action", "The elephant takes the jaggery. The summons stays on the straw."),
        ("blank", ""),
        ("trans", "FADE OUT."),
    ],
]


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _content_stream(items: list[tuple[str, str]]) -> bytes:
    ops = ["BT", "/F1 12 Tf"]
    y = TOP_Y
    for kind, text in items:
        if kind != "blank":
            ops.append(f"1 0 0 1 {X[kind]} {y} Tm ({_pdf_escape(text)}) Tj")
        y -= LEADING
    ops.append("ET")
    return "\n".join(ops).encode("latin-1")


def build_pdf(path: Path, pages: list[list[tuple[str, str]]] = VERI_PAGES) -> None:
    """Write a minimal valid PDF (Courier, one content stream per page) with correct xref offsets."""
    objects: list[bytes] = []
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, items in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode())
        stream = _content_stream(items)
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier /Encoding /WinAnsiEncoding >>")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


# --- ardha_song.docx -------------------------------------------------------------------------

ARDHA_SCENE_1: list[str] = [
    "ARDHA - song scene (golden test fixture)",
    "",
    "EXT. PERIYAR RIVERBANK - DUSK",
    "",
    "The river the colour of old brass. MEERA (24) sits on the ghat steps with a wet umbrella. Across the water a temple LOUDSPEAKER warms up.",
    "",
    "MEERA",
    "ഇന്നും വരില്ലെന്ന് അറിയാം. എന്നിട്ടും ഞാൻ ഇവിടെ ഇരിക്കുന്നു.",
    "",
    "VISHNU (O.S.)",
    "Meera! Vann nokk, ee side-il ninnu kaanam!",
    "",
    "She does not turn. She smiles.",
    "",
    'SONG: "ARDHA NILAVIL" - MONTAGE TO MUSIC',
    "",
    "♪ അർദ്ധ നിലാവിൽ ഞാൻ കാത്തിരുന്നു",
    "♪ Ardha nilaavil njan kaathirunnu",
    "♪ പുഴയുടെ കരയിൽ കണ്ണും നട്ട്",
    "♪ Nee varum ennu, nee varum ennu",
    "",
    "-- Meera walks the riverbank; the umbrella closes.",
    "-- Vishnu on the other bank, waving both arms.",
    "-- The loudspeaker song and the sung song fight for the air.",
]
ARDHA_SCENE_2: list[str] = [
    "INT. MEERA'S HOUSE - NIGHT",
    "",
    "A single bulb. AMMA (55) serves rice without looking up.",
    "",
    "AMMA",
    "പാട്ട് പാടി നടന്നാൽ വയറ് നിറയില്ല.",
    "",
    "MEERA",
    "(smiling)",
    "Amma, innu vayaru niranju.",
]


def build_docx(path: Path) -> None:
    """Plain Normal-style bilingual DOCX: blank paragraphs are the structure; a page break precedes scene 2."""
    from docx import Document

    document = Document()
    for line in ARDHA_SCENE_1:
        document.add_paragraph(line)
    document.add_page_break()
    for line in ARDHA_SCENE_2:
        document.add_paragraph(line)
    document.save(str(path))


if __name__ == "__main__":
    build_pdf(HERE / "veri_scenes.pdf")
    build_docx(HERE / "ardha_song.docx")
    print("wrote", HERE / "veri_scenes.pdf", "and", HERE / "ardha_song.docx")
