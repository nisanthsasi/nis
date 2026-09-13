"""Exports: breakdown/shot-list/hooks xlsx, Obsidian markdown vault, director's-notes PDF, hooks.md."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pdfplumber
import pytest
import yaml

from the_panel.export import (
    HAVE_REPORTLAB,
    export_breakdown_xlsx,
    export_directors_notes_pdf,
    export_hooks,
    export_hooks_xlsx,
    export_shotlist_xlsx,
    hooks_markdown,
    safe_filename,
    scene_markdown,
    winansi_safe,
    write_vault,
)
from the_panel.export.markdown import FILM_BRIEF_NOTE, INDEX_NOTE, STYLE_BIBLE_NOTE, bullets, group_by_sequence
from the_panel.export.pdf import layout_directors_notes
from the_panel.export.xlsx import BREAKDOWN_COLUMNS, HOOKS_COLUMNS, SET_PIECE_COLUMNS, SHOTLIST_COLUMNS
from the_panel.schemas import (
    Amendment,
    BreakdownRow,
    CbfcFlag,
    CinematographyBody,
    CinematographyOutput,
    CommercialHooksLedger,
    CostDriver,
    CostFlag,
    DecisionLogEntry,
    Department,
    DepartmentDirective,
    FilmBreakdown,
    HookEntry,
    MusicOutput,
    ProductionFlag,
    SetPieceCosting,
)
from the_panel.schemas.directive import ShotFrameRate
from the_panel.testing import build_example

from .conftest import make_film_brief, make_plan, make_scene, make_style_bible

MALAYALAM = "അഞ്ജുവിന്റെ കാൽക്കീഴിലെ നിലം പോയി"


def _cine(plan, **body_overrides) -> DepartmentDirective:
    out = build_example(CinematographyOutput)
    body = out.directive.model_copy(update=body_overrides)
    out = out.model_copy(update={"directive": body})
    return DepartmentDirective.from_output(Department.CINEMATOGRAPHY, plan.scene_id, plan.governing_idea, plan.pleasure_beat, out)


def _music(plan) -> DepartmentDirective:
    return DepartmentDirective.from_output(Department.MUSIC, plan.scene_id, plan.governing_idea, plan.pleasure_beat, build_example(MusicOutput))


def _sheet_rows(path: Path, sheet: str | None = None) -> list[list]:
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    return [[c.value for c in r] for r in ws.iter_rows()]


# ------------------------------------------------------------------ xlsx


def test_breakdown_xlsx_one_row_per_scene_with_set_piece_costing(tmp_path: Path):
    rows = [
        BreakdownRow(scene_id="S1", cast_speaking=["ANJU", "RAVI"], props=["PHONE", "ENVELOPE"], production_flags=[ProductionFlag.NIGHT_EXT, ProductionFlag.RAIN], cost_flag=CostFlag.MEDIUM, estimated_setups=6),
        BreakdownRow(scene_id="S3", cast_speaking=["ANJU"], background_count=300, production_flags=[ProductionFlag.CROWD, ProductionFlag.STUNT], cost_flag=CostFlag.HIGH, set_piece_costing=SetPieceCosting(extra_shoot_days=2, rehearsal_days=3, crowd_count=300, high_speed_days=1, star_windows=["Mar 3-5"], cheapest_staging_keeping_must_remember="one lane, 40 extras, long lens")),
    ]
    film = FilmBreakdown(budget_tier="low", budget_inr=40_000_000, night_ext_count=4, crowd_scene_ids=["S3"], feasibility_facts=["no crane", "S3 needs 300 extras"], cost_drivers=[CostDriver(driver="the barrier crowd", scene_ids=["S3"], cheaper_alternative="40 extras and a long lens")], cbfc_flags=[CbfcFlag(scene_id="S3", category="violence", likely_band="U/A", safer_staging="off-screen hit", dramatic_cost="loses the impact beat")])
    path = export_breakdown_xlsx(rows, film, tmp_path / "exports" / "breakdown.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Scenes"]
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold and ws.auto_filter.ref.startswith("A1:")
    grid = _sheet_rows(path, "Scenes")
    assert grid[0] == list(BREAKDOWN_COLUMNS + SET_PIECE_COLUMNS)
    by_id = {r[0]: dict(zip(grid[0], r)) for r in grid[1:]}
    assert by_id["S1"]["production_flags"] == "night_ext, rain" and by_id["S1"]["props"] == "PHONE, ENVELOPE" and by_id["S1"]["cost_flag"] == "MEDIUM"
    assert by_id["S1"]["set_piece"] is False and by_id["S1"]["sp_extra_shoot_days"] in (None, "")
    assert by_id["S3"]["set_piece"] is True and by_id["S3"]["sp_crowd_count"] == 300 and by_id["S3"]["sp_star_windows"] == "Mar 3-5"
    assert by_id["S3"]["sp_cheapest_staging_keeping_must_remember"] == "one lane, 40 extras, long lens"
    film_sheet = wb["Film"]
    titles = [r[0].value for r in film_sheet.iter_rows() if r[0].font.bold and r[0].value]
    assert titles[:2] == ["Film facts", "fact"] and "Cost drivers" in titles and "CBFC / OTT flags" in titles
    flat = [c.value for r in film_sheet.iter_rows() for c in r if c.value not in (None, "")]
    assert "no crane" in flat and "the barrier crowd" in flat and "off-screen hit" in flat and 40_000_000 in flat


def test_shotlist_xlsx_prefers_cinematography_directive(tmp_path: Path):
    p1, p2 = make_plan("S1", chosen="B"), make_plan("S2", chosen=None, recommended="A")
    cine = _cine(p1, slow_motion_frame_rates=[ShotFrameRate(shot_no=1, fps=96, earned_by="beat 2")], base_frame_rate=24)
    body = CinematographyBody.model_validate(cine.directive)
    body.shot_list[0].previz_prompt = "Anju in the doorway, sodium key from frame left, 24mm, slow push"
    body.shot_list[0].filtration = "1/4 Black Pro-Mist"
    body.shot_list[0].frame_rate = None  # the slow-motion table supplies it
    body.shot_list.append(body.shot_list[0].model_copy(update={"no": 2, "frame_rate": 48}))
    body.shot_list.append(body.shot_list[0].model_copy(update={"no": 3, "frame_rate": None}))
    cine = cine.model_copy(update={"directive": body.model_dump(mode="json")})
    path = export_shotlist_xlsx([p1, p2], {"S1": [cine, _music(p1)]}, tmp_path / "shotlist.xlsx")
    grid = _sheet_rows(path)
    assert grid[0] == list(SHOTLIST_COLUMNS)
    rows = [dict(zip(grid[0], r)) for r in grid[1:]]
    s1 = [r for r in rows if r["scene_id"] == "S1"]
    s2 = [r for r in rows if r["scene_id"] == "S2"]
    assert len(s1) == len(body.shot_list) and s1[0]["source"] == "cinematography" and s1[0]["option"] == "B"
    assert s1[0]["frame_rate"] == 96 and s1[0]["filtration"] == "1/4 Black Pro-Mist" and "sodium key" in s1[0]["previz_prompt"]
    assert s1[1]["frame_rate"] == 48 and s1[2]["frame_rate"] == 24  # a shot's own rate wins; otherwise the base rate
    assert s1[0]["aspect_ratio"] == body.shot_list[0].aspect_ratio
    assert s1[0]["cannot_live_without"] == (body.shot_list[0].no == p1.the_shot_it_cannot_live_without.shot_no)
    assert len(s2) == len(p2.chosen_option().shots) and s2[0]["source"] == "plan option A" and s2[0]["option"] == "A" and s2[0]["previz_prompt"] in (None, "")
    assert s2[0]["beat_ref"] == p2.chosen_option().shots[0].beat_ref and s2[0]["lens_mm"] == p2.chosen_option().shots[0].lens_mm
    assert openpyxl.load_workbook(path).active.freeze_panes == "A2"


def _ledger() -> CommercialHooksLedger:
    return CommercialHooksLedger(
        trailer_shots=[HookEntry(scene_id="S3", description="the barrier oner", kind="oner"), HookEntry(scene_id="S5", description="the hospital corridor walk")],
        poster_frames=[HookEntry(scene_id="S3", description="Anju under the sodium lamp")],
        bgm_hooks=[HookEntry(scene_id="S5", description="theme, first full statement", kind="theme")],
        interval_block=HookEntry(scene_id="S3", description="the barrier reversal", kind="reversal"),
        approved_scene_ids=["S3", "S5"],
    )


def test_hooks_xlsx_and_markdown(tmp_path: Path):
    ledger = _ledger()
    path = export_hooks_xlsx(ledger, tmp_path / "hooks.xlsx")
    grid = _sheet_rows(path, "Hooks")
    assert grid[0] == list(HOOKS_COLUMNS)
    assert grid[1] == ["interval_block", "S3", "reversal", "the barrier reversal"]
    assert ["trailer_shots", "S3", "oner", "the barrier oner"] in grid and ["bgm_hooks", "S5", "theme", "theme, first full statement"] in grid
    assert _sheet_rows(path, "Approved scenes") == [["scene_id"], ["S3"], ["S5"]]

    md = hooks_markdown(ledger)
    assert md.startswith("# Commercial Hooks Ledger") and "2 approved scenes: S3, S5" in md
    assert md.index("## Interval block") < md.index("## Trailer shots") < md.index("## Poster frames") < md.index("## Shareable moments")
    assert "- **S3** — the barrier oner _(oner)_" in md and "- **S5** — the hospital corridor walk" in md
    assert md.split("## Song slots")[1].split("##")[0].strip() == "- _none yet_"
    empty = hooks_markdown(CommercialHooksLedger())
    assert "no approved scenes yet" in empty and "- _not yet designated_" in empty

    written = export_hooks(ledger, tmp_path / "out")
    assert written["hooks.md"].read_text(encoding="utf-8") == md and openpyxl.load_workbook(written["hooks.xlsx"])["Hooks"].max_row == 6


# ------------------------------------------------------------------ markdown


def _frontmatter(md: str) -> dict:
    assert md.startswith("---\n")
    return yaml.safe_load(md.split("---\n", 2)[1])


def test_scene_markdown_frontmatter_sections_and_wikilinks():
    scene = make_scene("S3", 3, set_piece=True, must_remember="the barrier", load_bearing=True, must_feel=MALAYALAM, energy_target=8)
    plan = make_plan("S3", chosen=None)
    directives = [_cine(plan), _music(plan)]
    decision = DecisionLogEntry(scene_id="S3", chosen="B", human_note="B is braver", overrides=["no crane"], pushback_given=True, pushback_reason="the crane is the one move", style_bible_amendment=Amendment(scene_id="S3", change="allow the crane", reason="the barrier", affected_scene_ids=["S4"]), rerun_scene_ids=["S4"])
    md = scene_markdown(scene, plan, directives, ["device UNEARNED: slow motion → hold the wide"], decision, prev_id="S2", next_id="S4")
    fm = _frontmatter(md)
    assert fm == {
        "scene_id": "S3",
        "number": 3,
        "sequence": "SEQ1",
        "slug": "INT. KOCHI FLAT - NIGHT",
        "rasa_primary": "karuna",
        "rasa_secondary": "raudra",
        "must_feel": MALAYALAM,
        "pleasure_type": "tears",
        "set_piece": True,
        "must_remember": "the barrier",
        "energy_target": 8,
        "load_bearing": True,
        "option_chosen": "B",
        "cost_flag": "LOW",
        "tags": ["panel", "seq/SEQ1"],
    }
    for heading in ("## Must-feel", "## Governing idea", "## Pleasure beat", "## Director's note", "## Option A", "## Option B", "## The shot it cannot live without", "## Department stances", "## Directives", "## Audit flags", "## Decision", "## Links"):
        assert heading in md, heading
    assert "[[S2]]" in md and "[[S4]]" in md and "[[SEQ1]]" in md
    assert md.count("| # | Size | Angle |") == 2 and f"| 1 | {plan.option_A.shots[0].size.value} |" in md
    assert f"**Shot {plan.the_shot_it_cannot_live_without.shot_no}** — {plan.the_shot_it_cannot_live_without.why}" in md
    assert "<details><summary><b>cinematography</b>" in md and "<details><summary><b>music</b>" in md and md.count("</details>") == 2
    assert "- **shot_list:**" in md and "  - **1.**" in md and "    - **previz_prompt:**" in md and "- **DEPT_FLAG:**" in md
    assert "- device UNEARNED: slow motion → hold the wide" in md
    assert "**Chosen:** B" in md and "**Push-back given:** the crane is the one move" in md and "**Style Bible amendment (pending):** allow the crane" in md and "**Re-run:** [[S4]]" in md
    assert "**Colour:** " + plan.department_handoff.colour_stance in md


def test_scene_markdown_without_plan_or_decision():
    md = scene_markdown(make_scene("S1"), None, [], [], None)
    fm = _frontmatter(md)
    assert fm["option_chosen"] is None and fm["set_piece"] is False
    assert "_not yet integrated_" in md and "_no department directives yet_" in md and "_pending human review_" in md and "_none_" in md
    assert "## Links\nsequence [[SEQ1]]" in md and "[[S0]]" not in md
    assert _frontmatter(scene_markdown(make_scene("S9", sequence_id=None), None, [], [], None))["tags"] == ["panel", "seq/UNSEQUENCED"]


def test_bullets_drop_empty_values_but_keep_false_and_zero():
    out = bullets({"a": "x", "b": "", "c": None, "d": [], "e": False, "f": 0, "g": {"h": [1, {"no": 2, "k": "v"}]}})
    assert out == ["- **a:** x", "- **e:** False", "- **f:** 0", "- **g:**", "  - **h:**", "    - 1", "    - **2.**", "      - **no:** 2", "      - **k:** v"]


def test_safe_filename():
    assert safe_filename('INT. KOCHI FLAT / NIGHT: "take 2"?') == "INT. KOCHI FLAT - NIGHT- -take 2"
    assert safe_filename("S3") == "S3" and safe_filename("///") == "note" and safe_filename("[[S1]]#tag") == "S1-tag"


def test_write_vault_produces_expected_files(tmp_path: Path):
    brief, bible = make_film_brief(), make_style_bible()
    scenes = [make_scene("S1", 1), make_scene("S2", 2, must_feel=MALAYALAM), make_scene("S3", 3, set_piece=True, must_remember="the barrier"), make_scene("S4", 4, sequence_id="SEQ2"), make_scene("S7", 7, sequence_id=None)]
    plans = {"S1": make_plan("S1"), "S3": make_plan("S3")}
    directives = {"S1": [_cine(plans["S1"])]}
    decisions = {"S1": DecisionLogEntry(scene_id="S1", chosen="A")}
    root = tmp_path / "vault"
    written = write_vault(root, brief, bible, scenes, plans, directives, {"S3": ["engagement pleasure_gap: 12 min → move the laugh"]}, decisions)
    rel = sorted(p.relative_to(root).as_posix() for p in written)
    assert rel == sorted([FILM_BRIEF_NOTE, STYLE_BIBLE_NOTE, INDEX_NOTE, "scenes/S1.md", "scenes/S2.md", "scenes/S3.md", "scenes/S4.md", "scenes/S7.md", "sequences/SEQ1.md", "sequences/SEQ2.md", "sequences/UNSEQUENCED.md"])
    assert all(p.exists() for p in written)

    s2 = (root / "scenes/S2.md").read_text(encoding="utf-8")
    assert "← [[S1]]" in s2 and "[[S3]] →" in s2 and "sequence [[SEQ1]]" in s2 and MALAYALAM in s2
    s1 = (root / "scenes/S1.md").read_text(encoding="utf-8")
    assert "← [[" not in s1 and "[[S2]] →" in s1 and _frontmatter(s1)["option_chosen"] == "A"
    s3 = (root / "scenes/S3.md").read_text(encoding="utf-8")
    assert "- engagement pleasure_gap: 12 min → move the laugh" in s3 and _frontmatter(s3)["set_piece"] is True
    s7 = (root / "scenes/S7.md").read_text(encoding="utf-8")
    assert "← [[S4]]" in s7 and "sequence [[UNSEQUENCED]]" in s7 and _frontmatter(s7)["sequence"] == "UNSEQUENCED"

    seq1 = (root / "sequences/SEQ1.md").read_text(encoding="utf-8")
    assert _frontmatter(seq1) == {"sequence_id": "SEQ1", "function": "the smuggling begins", "scene_ids": ["S1", "S2", "S3"], "pleasure_type": "thrill", "set_piece_scene_id": "S3", "tags": ["panel", "sequence"]}
    assert "| [[S1]] |" in seq1 and "| [[S3]] |" in seq1 and "| yes |" in seq1
    index = (root / INDEX_NOTE).read_text(encoding="utf-8")
    assert "[[00-Film-Brief|Film Brief]]" in index and "[[01-Style-Bible|Style Bible]]" in index and "## [[SEQ1]] — the smuggling begins" in index
    assert "_5 scenes · 2 integrated · 1 decided_" in index and "· chosen A" in index and "set-piece" in index
    bible_md = (root / STYLE_BIBLE_NOTE).read_text(encoding="utf-8")
    assert _frontmatter(bible_md)["commercial_posture"] == 5 and "**5/10 (hybrid)**" in bible_md
    for needle in ("## Caps", "- elevation cues: 3", "## Device permissions", "- **slow motion** must pay off", "## Pleasure map", "[[SEQ1]] → thrill @ [[S3]]", "## Refusals", "- coverage-by-default", "## Lens affinity", "- primary: inarritu, ray_adoor, malayalam_new_wave", "## Amendments"):
        assert needle in bible_md, needle
    brief_md = (root / FILM_BRIEF_NOTE).read_text(encoding="utf-8")
    assert "# KAITHA — Film Brief" in brief_md and "- **inciting** → [[S1]]" in brief_md and "| ANJU |" in brief_md and "- [[S3]] — action (thrill) · INTERVAL" in brief_md


def test_group_by_sequence_orders_by_brief_then_number():
    brief = make_film_brief()
    scenes = [make_scene("S3", 3), make_scene("S1", 1), make_scene("S5", 5, sequence_id="SEQ2"), make_scene("S9", 9, sequence_id="SEQ9")]
    groups = group_by_sequence(brief, scenes)
    assert [(g[0], [s.id for s in g[2]]) for g in groups] == [("SEQ1", ["S1", "S3"]), ("SEQ2", ["S5"]), ("SEQ9", ["S9"])]


# ------------------------------------------------------------------ pdf


def _page_count(data: bytes) -> int:
    return len(re.findall(rb"/Type /Page(?!s)", data))


def test_directors_notes_pdf_one_page_per_scene(tmp_path: Path):
    scenes = [make_scene("S1", 1), make_scene("S2", 2), make_scene("S3", 3, set_piece=True, must_remember="the barrier")]
    plans = {"S1": make_plan("S1"), "S3": make_plan("S3")}
    path = export_directors_notes_pdf("KAITHA", scenes, plans, tmp_path / "exports" / "notes.pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF-1.4") and data.rstrip().endswith(b"%%EOF") and _page_count(data) == 3
    with pdfplumber.open(path) as doc:
        assert len(doc.pages) == 3 and doc.metadata.get("Title") == "KAITHA - Director's notes"
        p1 = doc.pages[0].extract_text()
        assert "S1 - INT. KOCHI FLAT - NIGHT" in p1 and "the floor has gone from under Anju" in p1
        assert plans["S1"].governing_idea in p1 and "Option A shot list" in p1 and f"Shot {plans['S1'].the_shot_it_cannot_live_without.shot_no}" in p1
        assert f"[beat {plans['S1'].option_A.shots[0].beat_ref}]" in p1
        p2 = doc.pages[1].extract_text()
        assert "S2 - " in p2 and "not yet integrated" in p2
        assert "SET-PIECE" in doc.pages[2].extract_text()


def test_pdf_survives_malayalam_with_documented_fallback(tmp_path: Path):
    scene = make_scene("S1", must_feel=MALAYALAM, slug="INT. വീട് - NIGHT")
    plan = make_plan("S1", directors_note=f"{MALAYALAM} — say it as a whisper")
    path = export_directors_notes_pdf("കൈത", [scene], {"S1": plan}, tmp_path / "ml.pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and _page_count(data) == 1
    with pdfplumber.open(path) as doc:
        text = doc.pages[0].extract_text()
    assert "see vault for original" in text and "? — say it as a whisper" in text  # em dash survives (WinAnsi), Malayalam collapses
    assert "INT. ? - NIGHT" in text and "? - Director's notes" in text
    assert winansi_safe("café — ‘quoted’ …") == ("café — ‘quoted’ …", False)
    assert winansi_safe(MALAYALAM) == ("? ? ? ?", True) and winansi_safe("") == ("", False) and winansi_safe("a\u2603b") == ("a?b", True)
    clean = export_directors_notes_pdf("KAITHA", [make_scene("S1")], {"S1": make_plan("S1")}, tmp_path / "clean.pdf")
    with pdfplumber.open(clean) as doc:
        assert "see vault for original" not in doc.pages[0].extract_text()


def test_pdf_overflow_continues_on_a_new_page(tmp_path: Path):
    long_note = " ".join(f"beat {i}: hold the look, do not rescue the silence." for i in range(160))
    scenes = [make_scene("S1", 1), make_scene("S2", 2)]
    plans = {"S1": make_plan("S1", directors_note=long_note), "S2": make_plan("S2")}
    layout = layout_directors_notes("KAITHA", scenes, plans)
    assert len(layout.pages) >= 3 and all(len(ln.text) < 200 for p in layout.pages for ln in p.lines)
    path = export_directors_notes_pdf("KAITHA", scenes, plans, tmp_path / "long.pdf", engine="minimal")
    assert _page_count(path.read_bytes()) == len(layout.pages)
    with pdfplumber.open(path) as doc:
        assert "S2 - INT. KOCHI FLAT - NIGHT" in doc.pages[-1].extract_text()


def test_pdf_engine_selection(tmp_path: Path):
    if HAVE_REPORTLAB:
        pytest.skip("reportlab installed: the fallback engine is exercised via engine='minimal' elsewhere")
    with pytest.raises(RuntimeError):
        export_directors_notes_pdf("KAITHA", [make_scene()], {}, tmp_path / "rl.pdf", engine="reportlab")
    auto = export_directors_notes_pdf("KAITHA", [make_scene()], {}, tmp_path / "auto.pdf")
    assert auto.read_bytes().startswith(b"%PDF-1.4")
